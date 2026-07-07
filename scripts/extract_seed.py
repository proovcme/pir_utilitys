from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SOURCE = Path("/Users/chernetchenko/Downloads/Калькулятор_оценки_конструктор_ПД_РД_ФЗИП.xlsx")
TARGET = Path("src/data/seedCatalog.ts")


def value(cell: Any) -> Any:
    if cell is None:
        return None
    if isinstance(cell, str):
        text = cell.strip()
        return text if text else None
    return cell


def number(v: Any, default: float = 0) -> float:
    if v is None or v == "":
        return default
    return float(v)


def flag(v: Any) -> bool:
    return bool(int(v or 0))


wb = load_workbook(SOURCE, data_only=True, read_only=False)

control = wb["Пульт"]
constructor = wb["Конструктор"]
rates_sheet = wb["Средние стоимостные группы"]
rd_sheet = wb["Справочник РД"]
pp87_sheet = wb["ПП87"]
sets_sheet = wb["Наборы"]

project_input = {
    "projectType": "",
    "address": "",
    "customer": "",
    "area": number(control["B4"].value),
    "vatRate": number(control["B5"].value),
    "commercialCoefficient": number(control["B6"].value),
    "overheadRate": 0.15,
    "bufferRate": number(control["B7"].value),
    "rateMultiplier": 1,
    "useGlobalCoefficient": False,
    "globalCoefficient": 1,
    "useGlobalDuration": False,
    "globalDurationDays": 30,
    "computerCost": 150000,
    "computerSalvageValue": 0,
    "computerUsefulLifeYears": 3,
    "insuranceContributionRate": 0.302,
    "includeCommon": flag(control["B8"].value),
    "presetPdOks": flag(control["B9"].value),
    "presetPdLinear": flag(control["B10"].value),
    "presetRdFull": flag(control["B11"].value),
    "presetRdCore": flag(control["B12"].value),
    "presetRdFrequent": flag(control["B13"].value),
}

line_map = {
    "id": 1,
    "source": 2,
    "category": 3,
    "section": 4,
    "name": 5,
    "performer": 6,
    "departmentCode": 7,
    "calculationType": 8,
    "presetPdOks": 9,
    "presetPdLinear": 10,
    "presetRdFull": 11,
    "presetRdCore": 12,
    "presetRdFrequent": 13,
    "common": 14,
    "manualInclude": 16,
    "excluded": 17,
    "workUnits": 19,
    "durationDays": 20,
    "manualAmount": 22,
    "coefficient": 24,
    "comment": 26,
}

lines = []
for row in range(2, constructor.max_row + 1):
    row_id = value(constructor.cell(row, 1).value)
    if not row_id:
        continue
    item: dict[str, Any] = {}
    for key, col in line_map.items():
        raw = value(constructor.cell(row, col).value)
        if key in {
            "presetPdOks",
            "presetPdLinear",
            "presetRdFull",
            "presetRdCore",
            "presetRdFrequent",
            "common",
            "manualInclude",
            "excluded",
        }:
            item[key] = flag(raw)
        elif key in {"workUnits", "durationDays", "manualAmount", "coefficient"}:
            item[key] = number(raw, 1 if key == "coefficient" else 0)
        else:
            item[key] = raw
    lines.append(item)

rates = []
for row in range(2, rates_sheet.max_row + 1):
    code = value(rates_sheet.cell(row, 1).value)
    if not code:
        continue
    rates.append(
        {
            "code": code,
            "group": value(rates_sheet.cell(row, 2).value),
            "monthlySalaryMedian": number(rates_sheet.cell(row, 3).value),
            "comment": value(rates_sheet.cell(row, 4).value),
        }
    )

rd_reference = []
for row in range(2, rd_sheet.max_row + 1):
    mark = value(rd_sheet.cell(row, 2).value)
    if not mark:
        continue
    rd_reference.append(
        {
            "number": value(rd_sheet.cell(row, 1).value),
            "mark": mark,
            "name": value(rd_sheet.cell(row, 3).value),
            "departmentCode": value(rd_sheet.cell(row, 4).value),
            "isCore": flag(rd_sheet.cell(row, 5).value),
            "isFrequentPublicBuilding": flag(rd_sheet.cell(row, 6).value),
        }
    )

pp87_reference = []
for row in range(2, pp87_sheet.max_row + 1):
    section_type = value(pp87_sheet.cell(row, 1).value)
    name = value(pp87_sheet.cell(row, 4).value)
    if not section_type or not name:
        continue
    pp87_reference.append(
        {
            "type": section_type,
            "number": value(pp87_sheet.cell(row, 2).value),
            "mark": value(pp87_sheet.cell(row, 3).value),
            "name": name,
            "departmentCode": value(pp87_sheet.cell(row, 5).value),
            "note": value(pp87_sheet.cell(row, 6).value),
        }
    )

sets = []
for row in range(2, sets_sheet.max_row + 1):
    name = value(sets_sheet.cell(row, 1).value)
    if not name:
        continue
    sets.append(
        {
            "name": name,
            "controlCell": value(sets_sheet.cell(row, 2).value),
            "description": value(sets_sheet.cell(row, 3).value),
            "exclusionHint": value(sets_sheet.cell(row, 4).value),
        }
    )

payload = {
    "version": 1,
    "sourceWorkbook": SOURCE.name,
    "projectInput": project_input,
    "lines": lines,
    "rates": rates,
    "rdReference": rd_reference,
    "pp87Reference": pp87_reference,
    "presetSets": sets,
}

TARGET.parent.mkdir(parents=True, exist_ok=True)
TARGET.write_text(
    "import type { SeedCatalog } from \"../domain/types\";\n\n"
    "export const seedCatalog = "
    + json.dumps(payload, ensure_ascii=False, indent=2)
    + " satisfies SeedCatalog;\n",
    encoding="utf-8",
)

print(f"Wrote {TARGET} with {len(lines)} lines, {len(rates)} rates")
